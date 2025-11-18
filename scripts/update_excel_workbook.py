#!/usr/bin/env python3
import argparse
import os
import sys
import json
import time
from datetime import datetime
from typing import List, Dict, Optional
from xml.etree import ElementTree as ET


def _ensure_openpyxl() -> Optional[object]:
    """Ensure openpyxl is available and return the module, installing if needed."""
    try:
        import openpyxl  # type: ignore
        return openpyxl
    except ImportError:
        try:
            import subprocess
            subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl'])
            import openpyxl  # type: ignore
            return openpyxl
        except Exception:
            return None


def _categorize_test(classname: str, file_path: str) -> str:
    """Determine test category based on classname or file path."""
    lower = (classname or '') + ' ' + (file_path or '')
    lower = lower.lower()
    if 'performance' in lower or 'perf' in lower:
        return 'Performance'
    if 'vision' in lower:
        return 'Vision'
    if 'integration' in lower:
        return 'Integration'
    if 'system' in lower:
        return 'System'
    if 'forecast' in lower or 'traffic_forecast' in lower:
        return 'Forecasting'
    return 'Unit'


def parse_junit_xml(junit_path: str) -> List[Dict[str, str]]:
    """Parse JUnit XML and return list of test result rows for Excel."""
    rows: List[Dict[str, str]] = []
    if not junit_path or not os.path.exists(junit_path):
        return rows
    try:
        tree = ET.parse(junit_path)
        root = tree.getroot()
        now_date = datetime.utcnow().strftime('%Y-%m-%d')
        for case in root.iter('testcase'):
            name = case.get('name') or ''
            classname = case.get('classname') or ''
            file_path = case.get('file') or ''
            failures = list(case.iter('failure'))
            skipped = list(case.iter('skipped'))
            error = list(case.iter('error'))
            passed = not failures and not skipped and not error
            score = '1.0' if passed else '0.0'
            status = 'Pass' if passed else 'Fail'
            comments = ''
            if failures:
                comments = failures[0].get('message') or failures[0].text or 'Failure'
            elif error:
                comments = error[0].get('message') or error[0].text or 'Error'
            elif skipped:
                status = 'Fail'
                comments = skipped[0].get('message') or 'Skipped'
                score = '0.0'
            cat = _categorize_test(classname, file_path)
            test_id = f"{classname}.{name}" if classname else name
            rows.append({
                'Test ID': test_id,
                'Test Name': name,
                'Date Conducted': now_date,
                'Score (%)': score,
                'Pass/Fail Status': status,
                'Comments': comments,
                'Test Category': cat,
            })
        return rows
    except Exception:
        return rows


def load_achievements(json_path: Optional[str]) -> List[Dict[str, str]]:
    """Load achievements from a JSON file if provided."""
    if not json_path:
        return []
    if not os.path.exists(json_path):
        return []
    try:
        with open(json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        rows: List[Dict[str, str]] = []
        for item in data:
            rows.append({
                'Achievement ID': str(item.get('id', '')),
                'Achievement Name': str(item.get('name', '')),
                'Description': str(item.get('description', '')),
                'Date Earned': str(item.get('date', datetime.utcnow().strftime('%Y-%m-%d'))),
                'Points Value': str(item.get('points', '')),
                'Category': str(item.get('category', '')),
            })
        return rows
    except Exception:
        return []


def _append_rows_openpyxl(workbook_path: str,
                          sheet_name: str,
                          headers: List[str],
                          rows: List[Dict[str, str]],
                          table_name: Optional[str] = None) -> bool:
    """Append rows into a worksheet using openpyxl and update table range if present."""
    if not rows:
        return True
    openpyxl = _ensure_openpyxl()
    if openpyxl is None:
        return False
    try:
        wb = openpyxl.load_workbook(workbook_path)
        ws = wb[sheet_name]
        # Ensure headers match
        sheet_headers = [ws.cell(row=1, column=i + 1).value for i in range(len(headers))]
        if sheet_headers != headers:
            for i, h in enumerate(headers):
                ws.cell(row=1, column=i + 1, value=h)
        # Find next empty row
        next_row = ws.max_row + 1
        # Append rows
        for r in rows:
            for col_idx, h in enumerate(headers, start=1):
                ws.cell(row=next_row, column=col_idx, value=r.get(h, ''))
            next_row += 1
        # Update table ref if present
        if table_name and hasattr(ws, '_tables'):
            for t in list(ws._tables):  # type: ignore[attr-defined]
                if t.name == table_name:
                    t.ref = f"A1:{openpyxl.utils.get_column_letter(len(headers))}{ws.max_row}"
                    break
        wb.save(workbook_path)
        return True
    except Exception:
        return False


def update_workbook(workbook_path: str,
                    achievements_json: Optional[str],
                    junit_xml: Optional[str]) -> bool:
    """Update the Excel workbook with achievements and test results."""
    achievements = load_achievements(achievements_json)
    tests = parse_junit_xml(junit_xml) if junit_xml else []
    ok1 = _append_rows_openpyxl(
        workbook_path,
        'Unique Achievements',
        ['Achievement ID', 'Achievement Name', 'Description', 'Date Earned', 'Points Value', 'Category'],
        achievements,
        table_name='tblAchievements'
    )
    ok2 = _append_rows_openpyxl(
        workbook_path,
        'Test Results',
        ['Test ID', 'Test Name', 'Date Conducted', 'Score (%)', 'Pass/Fail Status', 'Comments', 'Test Category'],
        tests,
        table_name='tblTestResults'
    )
    return ok1 and ok2


def main() -> None:
    """CLI entrypoint to update the Excel workbook from project outputs."""
    parser = argparse.ArgumentParser(description='Update Excel workbook with achievements and test results')
    parser.add_argument('--workbook-path', required=True, help='Path to AdaptiveTraffic_Workbook.xlsx')
    parser.add_argument('--achievements-json', required=False, help='Path to achievements JSON file')
    parser.add_argument('--junit-xml', required=False, help='Path to JUnit XML test results')
    args = parser.parse_args()

    start = time.time()
    ok = update_workbook(args.workbook_path, args.achievements_json, args.junit_xml)
    elapsed = time.time() - start
    if not ok:
        print(f"Update failed or partial. Time: {elapsed:.2f}s")
        sys.exit(1)
    print(f"Workbook updated successfully. Time: {elapsed:.2f}s")


if __name__ == '__main__':
    main()